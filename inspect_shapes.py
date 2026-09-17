import os, sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
base = r'C:\Users\bella\Downloads\Software Engineering-1\Software Engineering'

cases = [
    ('Lect_6', 'UML_Diagrams.pptx', [10, 11, 14, 15, 16]),
    ('Lect_7', 'Squences_ diagram.pptx', [4, 5, 6, 7, 8]),
    ('Lect_8', 'Activity diagram.pptx', [4, 5, 6, 7]),
    ('Lect_8', 'State chart diagram.pptx', [4, 5, 6])
]

for lect, fn, slides in cases:
    p = os.path.join(base, lect, fn)
    prs = Presentation(p)
    for s_idx in slides:
        slide = prs.slides[s_idx]
        print(f"=== {lect}/{fn} Slide {s_idx+1} ===")
        for s in slide.shapes:
            txt = s.text.strip().replace('\n', ' ') if s.has_text_frame else ''
            print(f"  Shape: {s.name} | Type: {s.shape_type} | Text: {txt[:60]}")
