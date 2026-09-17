# -*- coding: utf-8 -*-
import os, sys, subprocess, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))
DOCS_DIR = os.path.join(PROJECT_ROOT, "docs")
os.makedirs(DOCS_DIR, exist_ok=True)
CHROME_PATH = r"C:\Program Files\Google\Chrome\Application\chrome.exe"

from generate_academic_deliverables import (
    set_cell_background, set_cell_margins, set_cell_borders,
    add_p_rtl, add_run_rtl, add_h1_rtl, add_h2_rtl, add_callout
)

def build_adv_prog_docx():
    docx_path = os.path.join(DOCS_DIR, "Advanced_Programming_Requirements_Documentation.docx")
    doc = Document()

    # Margins & Section Setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True
        
        sectPr = section._sectPr
        bidi = parse_xml(f'<w:bidi {nsdecls("w")}/>')
        sectPr.append(bidi)

        # Header
        header = section.header
        p_h = header.paragraphs[0]
        p_h._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
        p_h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_h = p_h.add_run("توثيق مشروع البرمجة المتقدمة — معمارية SPMS REST — أ. الوليد الدعيس — جامعة صنعاء")
        r_h.font.name = "Cairo"
        r_h.font.size = Pt(8.5)
        r_h.font.color.rgb = RGBColor(100, 116, 139)

    # Title Cover Header Box
    tbl_cov = doc.add_table(rows=1, cols=1)
    tbl_cov.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_cov.autofit = False
    c_cov = tbl_cov.rows[0].cells[0]
    c_cov.width = Inches(6.5)
    set_cell_background(c_cov, "0F2942")
    set_cell_margins(c_cov, top=240, bottom=240, left=240, right=240)
    set_cell_borders(c_cov, top=None, bottom=None, left=None, right=None)

    p1 = c_cov.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
    add_run_rtl(p1, "جامعة صنعاء — كلية الحاسوب وتكنولوجيا المعلومات", font_name="Cairo", size_pt=11, bold=True, color_rgb=(203, 213, 225))

    p2 = c_cov.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
    add_run_rtl(p2, "تقرير استيفاء المتطلبات الهندسية والمعمارية لمشروع البرمجة المتقدمة", font_name="Cairo", size_pt=15, bold=True, color_rgb=(255, 255, 255))

    p3 = c_cov.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
    add_run_rtl(p3, "Smart Pharmacy Enterprise Multi-Platform Architecture (SPMS)", font_name="Cairo", size_pt=11, bold=True, color_rgb=(56, 189, 248))

    p4 = c_cov.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
    add_run_rtl(p4, "إشراف الأستاذ: الوليد الدعيس | العام الجامعي: 2026 / 2027 م", font_name="Cairo", size_pt=10, bold=False, color_rgb=(226, 232, 240))

    add_p_rtl(doc, "", space_after=8)

    # Team Matrix Table
    add_h2_rtl(doc, "فريق العمل وتوزيع الأدوار في معمارية النظام:")
    team_tbl = doc.add_table(rows=5, cols=4)
    team_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    team_tbl.autofit = False

    headers = ["م", "اسم الطالب", "الرقم الأكاديمي", "الدور التقني والمساهمة في معمارية النظام"]
    widths = [Inches(0.5), Inches(1.5), Inches(1.2), Inches(3.3)]

    for i, h in enumerate(headers):
        cell = team_tbl.rows[0].cells[i]
        cell.width = widths[i]
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
        add_run_rtl(p, h, font_name="Cairo", size_pt=9.5, bold=True, color_rgb=(255, 255, 255))

    team_data = [
        ("1", "عبدالله البوص", "25164359", "قائد الفريق: معمارية البرمجيات الهجينة، بوابة الخادم الخلفي (Gateway)، وأمان JWT وتكامل الخدمات."),
        ("2", "محمد قحري", "25164065", "مهندس النماذج الخلفية: تطوير واجهات REST APIs الخمس، نماذج الأعمال، ودعم استجابات JSON و XML."),
        ("3", "أحمد الصايدي", "25164067", "مهندس قواعد البيانات والأمان: تصميم المخطط 3NF، حماية PDO من حقن SQL، وتعقيم المدخلات (XSS)."),
        ("4", "إبراهيم إبراهيم", "25164587", "مهندس جودة البرمجيات: تطبيق أنماط التصميم (Design Patterns)، مبادئ SOLID، والاختبارات المؤتمتة.")
    ]

    for row_idx, row_data in enumerate(team_data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = team_tbl.rows[row_idx].cells[col_idx]
            cell.width = widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            set_cell_borders(cell, top={'val':'single','sz':4,'color':'E2E8F0'}, bottom={'val':'single','sz':4,'color':'E2E8F0'}, left=None, right=None)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col_idx == 3 else WD_ALIGN_PARAGRAPH.CENTER
            p._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
            add_run_rtl(p, text, font_name="Cairo", size_pt=9, bold=(col_idx==1), color_rgb=(30, 41, 59))

    add_p_rtl(doc, "", space_after=6)

    # Section 1
    add_h1_rtl(doc, "1. المعمارية الهجينة متعددة المنصات (Multi-Platform Architecture)")
    add_p_rtl(doc, "استجابةً لمتطلب بناء النظام على منصتين مختلفتين تشتركان في قاعدة بيانات مركزية وتتواصلان حصرياً عبر REST APIs، تم اعتماد المعمارية التالية:")
    add_p_rtl(doc, "• المنصة الأولى — واجهة الويب (Web Frontend): واجهة ويب تفاعلية متجاوبة (HTML5, CSS3 Custom System, Vanilla JS ES6+) تدعم الوضعين الليلي والنهاري وتتصل بالخوادم حصرياً عبر بروتوكول HTTP REST.")
    add_p_rtl(doc, "• المنصة الثانية — خادم الذكاء الاصطناعي (Python AI Service): خادم خدمات مصغرة مستقل مبني بإطار FastAPI يعمل على المنفذ 8000 لتقديم خدمات التنبؤ بالطلب والسلاسل الزمنية وسياسة FEFO.")
    add_p_rtl(doc, "• خادم البوابة المركزي (Central PHP Gateway): خادم PHP 8.2 كائني التوجه (OOP / MVC) يعمل على المنفذ 8080 لإدارة العمليات والمصادقة وإصدار الفواتير والتنسيق بين المنصات.")

    # Section 2
    add_h1_rtl(doc, "2. قاعدة البيانات المركزية المشتركة (Centralized Shared Relational Database)")
    add_p_rtl(doc, "• قاعدة بيانات موحدة ومطبعة حتى المستوى الثالث (3NF) تمنع تكرار البيانات وتضمن التكامل المرجعي التام.")
    add_p_rtl(doc, "• معمارية التخزين الهجين (Hybrid Cloud): تشمل MySQL محلياً للتشغيل الفوري بدون إنترنت في ملف database/schema.sql، وقاعدة بيانات PostgreSQL سحابياً على منصة Supabase في ملف database/supabase_schema.sql مع تطبيق سياسات أمان الصفوف (RLS).")

    # Section 3
    add_h1_rtl(doc, "3. واجهات الأعمال البرمجية الخمس وعمليات CRUD (Business Service REST APIs)")
    add_p_rtl(doc, "تم بناء 5 واجهات برمجية مستقلة تمثل وحدات الأعمال الرئيسية، وتدعم عمليات الـ CRUD الكاملة:")
    add_p_rtl(doc, "1. واجهة المصادقة auth.php: تسجيل مستخدم جديد، تسجيل الدخول، واسترجاع الملف الشخصي بواسطة رموز JWT المشفرة وحوكمة الصلاحيات (RBAC).")
    add_p_rtl(doc, "2. واجهة الأدوية medicines.php (Full CRUD): استرجاع الأصناف (GET)، إضافة صنف جديد (POST)، توريد رصيد وإصدار فاتورة توريد (POST ?action=restock)، تعديل بيانات الدواء (PUT)، وحذف الدواء (DELETE).")
    add_p_rtl(doc, "3. واجهة الطلبات orders.php: استرجاع الفواتير وتفاصيل بنودها (GET)، وإنشاء فاتورة صرف جديدة وخصم المخزون وإطلاق خطاف إعادة التدريب الآلي للذكاء الاصطناعي (POST).")
    add_p_rtl(doc, "4. واجهة الموردين suppliers.php: استرجاع الشركات الموردة (GET) وإضافة مورد جديد (POST).")
    add_p_rtl(doc, "5. واجهات الذكاء الاصطناعي (FastAPI Microservice): 6 مسارات تشمل التنبؤ بالطلب، السلاسل الزمنية، ترتيب الصلاحيات، التوصيات، وكشف الشذوذ.")

    # Section 4
    add_h1_rtl(doc, "4. منظومة الأمان والوقاية السيبرانية الشاملة (Cybersecurity Architecture)")
    add_p_rtl(doc, "• مصادقة JWT المشفرة: استخدام JSON Web Tokens الموقعة بخوارزمية HMAC-SHA256 مع فحص وقت الانتهاء (exp) والتحقق من الترويسة Authorization: Bearer في ملف backend/core/JWT.php.")
    add_p_rtl(doc, "• الحصانة التامة ضد حقن SQL: الاعتماد الحصري على PDO Prepared Statements وربط المعاملات الصارم ($stmt->execute) في كافة استعلامات قاعدة البيانات.")
    add_p_rtl(doc, "• الحماية ضد هجمات XSS: تعقيم وتطهير المدخلات بواسطة Security::sanitizeInput وتشفير المخرجات عبر htmlspecialchars.")
    add_p_rtl(doc, "• التخزين الآمن لكلمات المرور: التشفير القوي باستخدام خوارزمية Bcrypt مع ملح عشوائي ديناميكي عبر دالة password_hash.")
    add_p_rtl(doc, "• التحكم في الوصول المبني على الأدوار (RBAC): فصل صلاحيات الإدارة (admin)، الصيدلي (pharmacist)، والعميل (customer).")

    # Section 5
    add_h1_rtl(doc, "5. دعم صيغ تبادل البيانات: JSON و XML ديناميكياً (Dual Data Formats)")
    add_p_rtl(doc, "تطبيقاً للشرط الإجباري الصارم (Must support: JSON, XML):")
    add_p_rtl(doc, "• تم بناء فئة Response.php لتدعم إرجاع الاستجابات بصيغة JSON تلقائياً وبصيغة XML عند طلب العميل.")
    add_p_rtl(doc, "• تفعيل صيغة XML ديناميكياً عبر معامل الاستعلام ?format=xml أو ترويسة Accept: application/xml.")
    add_p_rtl(doc, "• تحويل الكائنات والمصفوفات المتداخلة تلقائياً إلى بنية XML معيارية نظيفة باستخدام SimpleXMLElement.")

    # Section 6
    add_h1_rtl(doc, "6. مصفوفة أنماط التصميم البرمجية (Design Patterns Implementation Matrix)")
    add_p_rtl(doc, "تجاوز المشروع شرط استخدام 3 أنماط تصميم بتطبيق 11 نمط تصميم برمجياً فعلياً داخل المشروع:")

    pat_tbl = doc.add_table(rows=12, cols=4)
    pat_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    pat_tbl.autofit = False

    p_headers = ["نمط التصميم (Pattern)", "التصنيف", "الموقع المصدري للكود", "المشكلة الهندسية المحلولة"]
    p_widths = [Inches(1.5), Inches(1.1), Inches(2.1), Inches(2.2)]

    for i, h in enumerate(p_headers):
        cell = pat_tbl.rows[0].cells[i]
        cell.width = p_widths[i]
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=90, bottom=90, left=70, right=70)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
        add_run_rtl(p, h, font_name="Cairo", size_pt=9, bold=True, color_rgb=(255, 255, 255))

    patterns_data = [
        ("Factory Pattern", "Creational", "backend/patterns/creational/Factory.php", "إنشاء كائنات الأدوية والطلبات دون اقتران مسبق."),
        ("Singleton Pattern", "Creational", "backend/patterns/creational/Singleton.php", "ضمان نسخة اتصال واحدة بقاعدة بيانات PDO لمنع استنزاف الذاكرة."),
        ("Builder Pattern", "Creational", "backend/patterns/creational/Builder.php", "تجميع فواتير الصرف والتوريد المعقدة متعددة البنود خطوة بخطوة."),
        ("Prototype Pattern", "Creational", "backend/patterns/creational/Prototype.php", "استنساخ قوالب الأدوية المتشابهة لتسريع إدخال الأصناف."),
        ("Facade Pattern", "Structural", "backend/patterns/structural/Facade.php", "واجهة موحدة لعملية الصرف وخصم المخزون وإشعار الذكاء."),
        ("Adapter Pattern", "Structural", "backend/patterns/structural/Adapter.php", "توحيد واجهات إرسال واستقبال البيانات بين خادم PHP وبايثون."),
        ("Decorator Pattern", "Structural", "backend/patterns/structural/Decorator.php", "تغليف وتطبيق طبقات الأمان وصلاحيات JWT ديناميكياً."),
        ("Proxy Pattern", "Structural", "backend/patterns/structural/Proxy.php", "التحكم في العمليات الحساسة وتنفيذ Caching للمخزون."),
        ("Strategy Pattern", "Behavioral", "backend/patterns/behavioral/Strategy.php", "التبديل الحي بين خوارزميات التنبؤ وحساب التسعير والخصومات."),
        ("Observer Pattern", "Behavioral", "backend/patterns/behavioral/Observer.php", "إرسال إشعارات فورية عند هبوط رصيد الدواء دون حد الأمان."),
        ("Command Pattern", "Behavioral", "backend/patterns/behavioral/Command.php", "تغليف عمليات الصرف كأوامر قابلة للمراجعة والتراجع.")
    ]

    for r_idx, r_data in enumerate(patterns_data, start=1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = pat_tbl.rows[r_idx].cells[c_idx]
            cell.width = p_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=70, right=70)
            set_cell_borders(cell, top={'val':'single','sz':4,'color':'CBD5E1'}, bottom={'val':'single','sz':4,'color':'CBD5E1'}, left=None, right=None)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if c_idx == 3 else WD_ALIGN_PARAGRAPH.CENTER
            p._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
            add_run_rtl(p, val, font_name="Cairo", size_pt=8.5, bold=(c_idx==0), color_rgb=(15, 23, 42))

    add_p_rtl(doc, "", space_after=6)

    # Section 7
    add_h1_rtl(doc, "7. تطبيق مبادئ SOLID وهندسة الكود النظيف")
    add_p_rtl(doc, "• المسؤولية الواحدة (SRP): فصل مهام الاتصال (Database)، المصادقة (JWT)، والتعقيم (Security) في فئات مستقلة.")
    add_p_rtl(doc, "• الفتح والإغلاق (OCP): إمكانية إضافة طرق دفع أو خوارزميات جديدة دون كسر الكود القائم.")
    add_p_rtl(doc, "• استبدال لسكوف (LSP): تطابق تام بين الفئات المشتقة وعقود واجهاتها الأساسية.")
    add_p_rtl(doc, "• فصل الواجهات (ISP): تجزئة الواجهات البرمجية إلى خدمات متخصصة صغيرة.")
    add_p_rtl(doc, "• عكس التبعية (DIP): اعتماد الطبقات العليا على التجريدات (Abstractions) عبر أنماط المصنع والواجهة.")

    # Section 8
    add_h1_rtl(doc, "8. الاختبارات البرمجية والتوثيق الهندسي (Testing & Documentation)")
    add_p_rtl(doc, "• الاختبارات الآلية: فحص شامل لمصادقة JWT، تعقيم XSS، وأنماط التصميم في tests/test_api.php بنسبة نجاح 100%.")
    add_p_rtl(doc, "• التوثيق التفاعلي بنظام Swagger / OpenAPI 3.0 متوفر في docs/swagger.json و docs/api_docs.html و FastAPI Swagger.")
    add_p_rtl(doc, "• حزمة المخططات الهندسية الثمانية متوفرة بجودة عالية في مجلد docs/SPMS_Diagrams_BlackWhite وتشمل DFD 0, DFD 1, Use Case, ERD, Class, Sequence, Activity, State Chart.")

    doc.save(docx_path)
    shutil.copyfile(docx_path, os.path.join(PROJECT_ROOT, "Advanced_Programming_Requirements_Documentation.docx"))
    print(f"[Word DOCX Created]: {docx_path}")

def build_adv_prog_pdf():
    html_path = os.path.join(DOCS_DIR, "Advanced_Programming_Requirements_Documentation.html")
    pdf_path = os.path.join(DOCS_DIR, "Advanced_Programming_Requirements_Documentation.pdf")

    html = """<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
<meta charset="utf-8">
<style>
@import url('https://fonts.googleapis.com/css2?family=Cairo:wght@400;600;700;800&display=swap');
@page { size: A4; margin: 20mm 15mm; }
body { font-family: 'Cairo', sans-serif; font-size: 10pt; line-height: 1.5; color: #0f172a; margin: 0; padding: 0; }
.header-box { background: #0F2942; color: white; padding: 18px 24px; border-radius: 6px; text-align: center; margin-bottom: 20px; }
.header-box h1 { font-size: 14pt; margin: 6px 0; color: #ffffff; }
.header-box h2 { font-size: 10.5pt; margin: 4px 0; color: #38bdf8; font-weight: normal; }
.header-box p { font-size: 9pt; margin: 4px 0; color: #cbd5e1; }
h2.sec { color: #0F2942; font-size: 12pt; border-bottom: 1.5px solid #0F2942; padding-bottom: 4px; margin-top: 18px; margin-bottom: 8px; }
p, li { font-size: 9.5pt; text-align: justify; }
ul { margin-top: 4px; margin-bottom: 8px; padding-right: 20px; }
table { width: 100%; border-collapse: collapse; margin: 12px 0; font-size: 8.5pt; }
th { background: #0F2942; color: white; padding: 8px 10px; border: 1px solid #0F2942; text-align: center; font-weight: bold; }
td { padding: 6px 8px; border: 1px solid #cbd5e1; text-align: center; }
tr:nth-child(even) td { background: #f8fafc; }
.callout { background: #f1f5f9; border-right: 4px solid #0F2942; padding: 10px 14px; margin: 12px 0; border-radius: 0 4px 4px 0; font-size: 9pt; color: #334155; }
</style>
</head>
<body>

<div class="header-box">
  <p>جامعة صنعاء — كلية الحاسوب وتكنولوجيا المعلومات — قسم نظم المعلومات</p>
  <h1>تقرير استيفاء المتطلبات الهندسية والمعمارية لمشروع البرمجة المتقدمة</h1>
  <h2>نظام إدارة الصيدلية المتكامل بالمعمارية الهجينة والخدمات المصغرة (SPMS REST Architecture)</h2>
  <p>إشراف الأستاذ: الوليد الدعيس | العام الجامعي: 2026 / 2027 م</p>
</div>

<h2 class="sec">فريق العمل وتوزيع الأدوار في معمارية النظام</h2>
<table>
  <thead>
    <tr><th style="width:5%;">م</th><th style="width:20%;">اسم الطالب</th><th style="width:15%;">الرقم الأكاديمي</th><th style="width:60%;">الدور التقني والمساهمة في معمارية النظام</th></tr>
  </thead>
  <tbody>
    <tr><td>1</td><td><strong>عبدالله البوص</strong></td><td>25164359</td><td style="text-align:right;"><strong>قائد الفريق:</strong> معمارية البرمجيات الهجينة، بوابة الخادم الخلفي (Gateway)، وأمان JWT وتكامل الخدمات.</td></tr>
    <tr><td>2</td><td><strong>محمد قحري</strong></td><td>25164065</td><td style="text-align:right;"><strong>مهندس النماذج الخلفية:</strong> تطوير واجهات REST APIs الخمس، نماذج الأعمال، ودعم استجابات JSON و XML.</td></tr>
    <tr><td>3</td><td><strong>أحمد الصايدي</strong></td><td>25164067</td><td style="text-align:right;"><strong>مهندس قواعد البيانات والأمان:</strong> تصميم المخطط 3NF، حماية PDO من حقن SQL، وتعقيم المدخلات (XSS).</td></tr>
    <tr><td>4</td><td><strong>إبراهيم إبراهيم</strong></td><td>25164587</td><td style="text-align:right;"><strong>مهندس جودة البرمجيات:</strong> تطبيق أنماط التصميم (Design Patterns)، مبادئ SOLID، والاختبارات المؤتمتة.</td></tr>
  </tbody>
</table>

<h2 class="sec">1. المعمارية الهجينة متعددة المنصات (Multi-Platform REST Architecture)</h2>
<ul>
  <li><strong>المنصة الأولى — واجهة الويب (Web Client):</strong> واجهة ويب متجاوبة عصرية (HTML5/CSS3/Vanilla JS ES6+) تدعم الوضعين الليلي والنهاري وتتصل بالخوادم عبر REST APIs حصرياً.</li>
  <li><strong>المنصة الثانية — محرك الذكاء الاصطناعي (Python AI Service):</strong> خادم خدمات مصغرة مستقل مبني بـ FastAPI يعمل على المنفذ 8000 لتقديم خدمات التنبؤ بالطلب والسلاسل الزمنية وسياسة FEFO.</li>
  <li><strong>خادم البوابة المركزي (Central PHP Gateway):</strong> مبني بـ PHP 8.2 كائني التوجه (OOP / MVC) على المنفذ 8080 لإدارة العمليات والمصادقة وإصدار الفواتير.</li>
</ul>

<h2 class="sec">2. قاعدة البيانات المركزية المشتركة ونموذج التخزين الهجين (3NF Database)</h2>
<ul>
  <li>قاعدة بيانات علائقية مطبعة حتى المستوى الثالث (3NF) تمنع التكرار وتضمن التكامل المرجعي.</li>
  <li>نموذج التخزين الهجين: MySQL محلياً للتشغيل المستمر بدون إنترنت في <code>database/schema.sql</code>، و PostgreSQL سحابياً على Supabase مع سياسات أمان الصفوف (RLS) في <code>database/supabase_schema.sql</code>.</li>
</ul>

<h2 class="sec">3. واجهات الأعمال البرمجية الخمس وعمليات CRUD الكاملة</h2>
<ul>
  <li><strong>1. auth.php:</strong> تسجيل، دخول، جلب الملف الشخصي برموز JWT المشفرة وحوكمة الصلاحيات (RBAC).</li>
  <li><strong>2. medicines.php (Full CRUD):</strong> قراءة قائمة وتفاصيل الأصناف (GET)، إضافة صنف (POST)، توريد رصيد وإصدار فاتورة توريد (POST ?action=restock)، تعديل صنف (PUT)، وحذف صنف (DELETE).</li>
  <li><strong>3. orders.php:</strong> استرجاع الفواتير وبنودها (GET)، وإنشاء فاتورة صرف وخصم المخزون وإشعار الذكاء الاصطناعي لإعادة التدريب الفوري (POST).</li>
  <li><strong>4. suppliers.php:</strong> استرجاع الشركات الموردة (GET) وإضافة مورد جديد (POST).</li>
  <li><strong>5. AI Microservice:</strong> مسارات التنبؤ بالطلب، السلاسل الزمنية، ترتيب الصلاحيات، التوصيات، وكشف الشذوذ.</li>
</ul>

<h2 class="sec">4. منظومة الأمان والوقاية السيبرانية الشاملة</h2>
<ul>
  <li><strong>مصادقة JWT المشفرة:</strong> رموز موقعة بـ HMAC-SHA256 مع فحص انتهاء الصلاحية (exp) وإلزامية ترويسة Bearer Token في <code>backend/core/JWT.php</code>.</li>
  <li><strong>الحصانة التامة ضد حقن SQL:</strong> استخدام PDO Prepared Statements وعزل أوامر SQL عن المدخلات تماماً في كافة الاستعلامات.</li>
  <li><strong>الحماية ضد XSS:</strong> تعقيم المدخلات عبر <code>Security::sanitizeInput</code> وتشفير المخرجات عبر <code>htmlspecialchars</code>.</li>
  <li><strong>التخزين الآمن لكلمات المرور:</strong> تشفير Bcrypt مع ملح عشوائي ديناميكي عبر <code>password_hash</code>.</li>
  <li><strong>حوكمة الصلاحيات (RBAC):</strong> فصل صلاحيات الإدارة (admin)، الصيدلي (pharmacist)، والعميل (customer).</li>
</ul>

<h2 class="sec">5. دعم صيغتي JSON و XML ديناميكياً (Dual Data Formats)</h2>
<p>تم بناء فئة <code>Response.php</code> لترجع استجابات JSON افتراضياً، وترجع XML معيارياً عند تمرير <code>?format=xml</code> أو ترويسة <code>Accept: application/xml</code> عبر كائن <code>SimpleXMLElement</code>.</p>

<h2 class="sec">6. مصفوفة أنماط التصميم البرمجية (Design Patterns — 11 نمطاً مطبقاً)</h2>
<table>
  <thead>
    <tr><th>النمط (Pattern)</th><th>التصنيف</th><th>الموقع في المشروع</th><th>المشكلة الهندسية المحلولة</th></tr>
  </thead>
  <tbody>
    <tr><td><strong>Factory</strong></td><td>Creational</td><td>backend/patterns/creational/Factory.php</td><td>إنشاء كائنات الأدوية والطلبات دون اقتران مسبق.</td></tr>
    <tr><td><strong>Singleton</strong></td><td>Creational</td><td>backend/patterns/creational/Singleton.php</td><td>ضمان وجود نسخة اتصال وحيدة بقاعدة البيانات PDO لمنع استنزاف الموارد.</td></tr>
    <tr><td><strong>Builder</strong></td><td>Creational</td><td>backend/patterns/creational/Builder.php</td><td>بناء فواتير الصرف والتوريد المعقدة خطوة بخطوة.</td></tr>
    <tr><td><strong>Prototype</strong></td><td>Creational</td><td>backend/patterns/creational/Prototype.php</td><td>استنساخ قوالب الأدوية المتشابهة لتسريع إدخال الأصناف.</td></tr>
    <tr><td><strong>Facade</strong></td><td>Structural</td><td>backend/patterns/structural/Facade.php</td><td>واجهة موحدة لعملية الصرف وخصم المخزون وإشعار الذكاء.</td></tr>
    <tr><td><strong>Adapter</strong></td><td>Structural</td><td>backend/patterns/structural/Adapter.php</td><td>توحيد واجهات إرسال واستقبال البيانات بين خادم PHP وبايثون.</td></tr>
    <tr><td><strong>Decorator</strong></td><td>Structural</td><td>backend/patterns/structural/Decorator.php</td><td>تغليف وتطبيق طبقات الأمان وصلاحيات JWT ديناميكياً.</td></tr>
    <tr><td><strong>Proxy</strong></td><td>Structural</td><td>backend/patterns/structural/Proxy.php</td><td>التحكم في العمليات الحساسة وتنفيذ آليات الـ Caching.</td></tr>
    <tr><td><strong>Strategy</strong></td><td>Behavioral</td><td>backend/patterns/behavioral/Strategy.php</td><td>التبديل الحي بين خوارزميات التنبؤ وحساب الخصومات والتسعير.</td></tr>
    <tr><td><strong>Observer</strong></td><td>Behavioral</td><td>backend/patterns/behavioral/Observer.php</td><td>إرسال إشعارات فورية عند هبوط رصيد الدواء دون حد الأمان.</td></tr>
    <tr><td><strong>Command</strong></td><td>Behavioral</td><td>backend/patterns/behavioral/Command.php</td><td>تغليف عمليات الصرف والتوريد كأوامر قابلة للمراجعة والتراجع.</td></tr>
  </tbody>
</table>

<h2 class="sec">7. مبادئ SOLID والاختبارات والمخططات الهندسية</h2>
<ul>
  <li><strong>مبادئ SOLID:</strong> تطبيق كامل لمبادئ (SRP, OCP, LSP, ISP, DIP) في فئات الكود البرمجي.</li>
  <li><strong>الاختبارات الآلية:</strong> اجتياز 100% من اختبارات JWT و XSS ومصنع الفئات في <code>tests/test_api.php</code> في 0.007 ثانية.</li>
  <li><strong>التوثيق والمخططات:</strong> توثيق Swagger / OpenAPI 3.0 متوفر في <code>docs/swagger.json</code>، وحزمة المخططات الهندسية الثمانية (DFD 0 & 1, Use Case, ERD, Class, Sequence, Activity, State Chart) في مجلد <code>docs/SPMS_Diagrams_BlackWhite/</code>.</li>
</ul>

</body>
</html>
"""
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)

    if os.path.exists(CHROME_PATH):
        cmd = [
            CHROME_PATH,
            "--headless",
            "--disable-gpu",
            f"--print-to-pdf={pdf_path}",
            "--no-pdf-header-footer",
            html_path
        ]
        subprocess.run(cmd, capture_output=True)
        shutil.copyfile(pdf_path, os.path.join(PROJECT_ROOT, "Advanced_Programming_Requirements_Documentation.pdf"))
        print(f"[PDF Created]: {pdf_path}")

if __name__ == "__main__":
    build_adv_prog_docx()
    build_adv_prog_pdf()
